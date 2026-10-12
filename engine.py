import os
import re
import json

import pdfplumber
import docx

import config

SECTION_WORDS = {
    "summary": ["summary", "objective", "profile", "about me", "career objective",
                "professional summary", "career summary", "professional profile"],
    "skills": ["skills", "technical skills", "key skills", "core competencies",
               "technologies", "professional skills", "skill set", "technical expertise"],
    "experience": ["experience", "work experience", "professional experience",
                   "employment", "work history", "career history", "internship",
                   "internships", "industrial training"],
    "education": ["education", "academic", "qualification", "academics",
                  "educational qualification", "academic qualification"],
    "projects": ["projects", "academic projects", "key projects", "personal projects",
                 "major projects", "project work"],
    "certifications": ["certification", "certifications", "courses", "training",
                       "achievements", "awards", "accomplishments",
                       "certifications and achievements"],
    "activities": ["activities", "extra curricular activities", "co curricular activities",
                   "volunteer experience", "positions of responsibility"],
    "languages": ["languages", "languages known"],
}

BULLET_CHARS = ("•", "●", "▪", "◦", "-", "*", "‣")


def load_data():
    path = os.path.join(config.BASE_DIR, "data.json")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def read_pdf(path):
    pages = []
    table_count = 0
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
            table_count = table_count + len(page.find_tables())
    return "\n".join(pages), table_count


def read_docx(path):
    doc = docx.Document(path)
    lines = [p.text for p in doc.paragraphs]

    # table content is often used for skills, so pick that up too
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines), len(doc.tables)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    return "", 0


def clean_text(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_contact(text):
    contact = {"email": "", "phone": "", "linkedin": "", "github": ""}

    m = re.search(r"[\w\.\-\+]+@[\w\-]+\.[\w\.\-]+", text)
    if m:
        contact["email"] = m.group(0)

    m = re.search(r"(\+91[\-\s]?)?[6-9]\d{9}", text)
    if m:
        contact["phone"] = m.group(0)

    m = re.search(r"(linkedin\.com/in/[\w\-]+)", text, re.I)
    if m:
        contact["linkedin"] = m.group(1)

    m = re.search(r"(github\.com/[\w\-]+)", text, re.I)
    if m:
        contact["github"] = m.group(1)

    return contact


def is_heading(line):
    line = line.strip()
    if not line or len(line) > 50:
        return False

    plain = re.sub(r"[^a-z& ]", " ", line.lower())
    plain = re.sub(r"\s+", " ", plain).strip()
    plain = plain.replace("&", "and").replace(" and ", " ")

    for name, words in SECTION_WORDS.items():
        if plain in words:
            return name
    return False


def split_sections(text):
    sections = {}
    current = None
    buffer = []

    for line in text.split("\n"):
        found = is_heading(line)
        if found:
            if current and current not in sections:
                sections[current] = "\n".join(buffer).strip()
            current = found
            buffer = []
        elif current:
            buffer.append(line)

    if current and current not in sections:
        sections[current] = "\n".join(buffer).strip()

    return sections


def normal_skills(data):
    skills = []
    for name, group in data["skills"].items():
        if name != "short_skills":
            skills.extend(group)
    return skills


def match_skill(skill, text):
    pattern = r"(?<![a-z0-9+#\.\-])" + re.escape(skill) + r"(?![a-z0-9+#\.\-])"
    return re.search(pattern, text) is not None


def extract_skills(text, data, skills_text=""):
    found = []
    low = text.lower()

    for skill in normal_skills(data):
        if match_skill(skill, low):
            found.append(skill)

    # single letter names like c, r, go give false matches in normal prose,
    # so look for them only inside the skills section
    if skills_text:
        low_skills = skills_text.lower()
        for skill in data["skills"].get("short_skills", []):
            if match_skill(skill, low_skills):
                found.append(skill)

    return sorted(set(found))


def get_bullets(text):
    bullets = []
    for line in text.split("\n"):
        line = line.strip()
        if line.startswith(BULLET_CHARS) and len(line) > 10:
            bullets.append(line)
    return bullets


def parse_resume(path):
    raw, table_count = extract_text(path)
    raw = clean_text(raw)
    data = load_data()

    sections = split_sections(raw)
    bullets = get_bullets(raw)

    parsed = {
        "contact": extract_contact(raw),
        "sections": list(sections.keys()),
        "skills": extract_skills(raw, data, sections.get("skills", "")),
        "word_count": len(raw.split()),
        "bullet_count": len(bullets),
        "quantified_bullets": len([b for b in bullets if re.search(r"\d", b)]),
        "table_count": table_count,
        "contact_at_top": bool(extract_contact("\n".join(raw.split("\n")[:12]))["email"]),
    }

    return raw, parsed


def is_fresher(experience_years):
    try:
        return float(experience_years or 0) < 1
    except (TypeError, ValueError):
        return True


def score_contact(parsed):
    points = 0
    notes = []
    contact = parsed["contact"]

    if contact["email"]:
        points = points + 4
    else:
        notes.append("Email address not found in the resume.")

    if contact["phone"]:
        points = points + 3
    else:
        notes.append("Phone number not found.")

    if contact["linkedin"] or contact["github"]:
        points = points + 3
    else:
        notes.append("Add a LinkedIn or GitHub profile link.")

    return points, notes


def score_sections(parsed, fresher):
    if fresher:
        needed = ["summary", "skills", "projects", "education"]
    else:
        needed = ["summary", "skills", "experience", "education"]

    found = parsed["sections"]
    points = 0
    missing = []

    for name in needed:
        if name in found:
            points = points + 5
        else:
            missing.append(name)

    notes = []
    if missing:
        notes.append("Missing sections: " + ", ".join(missing))

    return points, notes


def score_keywords(parsed, role_data):
    must = role_data["must_have"]
    good = role_data["good_to_have"]
    have = parsed["skills"]

    must_hit = [s for s in must if s in have]
    good_hit = [s for s in good if s in have]

    must_points = (len(must_hit) / len(must)) * 20 if must else 0
    good_points = (len(good_hit) / len(good)) * 10 if good else 0
    points = int(round(must_points + good_points))

    notes = []
    missing_must = [s for s in must if s not in have]
    if missing_must:
        notes.append("Must-have skills for this role were not found in the resume: " + ", ".join(missing_must))

    return points, notes


def score_format(parsed):
    points = 0
    notes = []

    if parsed["bullet_count"] >= 5:
        points = points + 7
    elif parsed["bullet_count"] > 0:
        points = points + 4
        notes.append("Too few bullet points. Keep 3 to 5 bullets under each role or project.")
    else:
        notes.append("No bullet points used. ATS software reads plain paragraphs poorly.")

    if parsed["table_count"] == 0:
        points = points + 7
    else:
        notes.append("A table was found. ATS software often skips tables, so use plain text instead.")

    if parsed["contact_at_top"]:
        points = points + 6
    else:
        notes.append("Contact details should appear at the top of the resume.")

    return points, notes


def score_impact(parsed):
    total = parsed["bullet_count"]
    quantified = parsed["quantified_bullets"]

    if total == 0:
        return 0, ["Write achievements as bullet points and include numbers."]

    ratio = quantified / total
    points = int(round(min(ratio / 0.4, 1) * 10))

    notes = []
    if ratio < 0.4:
        notes.append(
            "Only %d of %d bullets contain numbers. Lines like 'reduced load time by 40%%' show real impact."
            % (quantified, total)
        )

    return points, notes


def score_length(parsed):
    words = parsed["word_count"]

    if 350 <= words <= 900:
        return 10, []
    if words < 350:
        return 5, ["Resume is short (%d words). Add more detail to projects and skills." % words]
    if words <= 1200:
        return 7, ["Resume is a bit long (%d words). Keep it within two pages." % words]
    return 4, ["Resume is too long (%d words). Remove older and less relevant detail." % words]


def ats_score(parsed, role_data, experience_years):
    fresher = is_fresher(experience_years)

    contact_pts, contact_notes = score_contact(parsed)
    section_pts, section_notes = score_sections(parsed, fresher)
    keyword_pts, keyword_notes = score_keywords(parsed, role_data)
    format_pts, format_notes = score_format(parsed)
    impact_pts, impact_notes = score_impact(parsed)
    length_pts, length_notes = score_length(parsed)

    breakdown = [
        {"name": "Contact Details", "score": contact_pts, "max": 10},
        {"name": "Resume Sections", "score": section_pts, "max": 20},
        {"name": "Role Keyword Match", "score": keyword_pts, "max": 30},
        {"name": "ATS Formatting", "score": format_pts, "max": 20},
        {"name": "Quantified Impact", "score": impact_pts, "max": 10},
        {"name": "Resume Length", "score": length_pts, "max": 10},
    ]

    total = sum(b["score"] for b in breakdown)
    notes = contact_notes + section_notes + keyword_notes + format_notes + impact_notes + length_notes

    return {
        "total": total,
        "mode": "Fresher" if fresher else "Experienced",
        "breakdown": breakdown,
        "notes": notes,
    }


def skill_gap(parsed, role_data):
    have = parsed["skills"]

    must_missing = [s for s in role_data["must_have"] if s not in have]
    good_missing = [s for s in role_data["good_to_have"] if s not in have]

    total_weight = len(role_data["must_have"]) * 2 + len(role_data["good_to_have"])
    have_weight = (len(role_data["must_have"]) - len(must_missing)) * 2 + \
                  (len(role_data["good_to_have"]) - len(good_missing))

    match_percent = int(round((have_weight / total_weight) * 100)) if total_weight else 0

    return {
        "match_percent": match_percent,
        "matched": [s for s in role_data["must_have"] + role_data["good_to_have"] if s in have],
        "must_missing": must_missing,
        "good_missing": good_missing,
    }


def suggest_courses(gap, data):
    picks = []
    for skill in gap["must_missing"] + gap["good_missing"]:
        for course in data["courses"].get(skill, []):
            item = dict(course)
            item["skill"] = skill
            item["priority"] = "Must have" if skill in gap["must_missing"] else "Good to have"
            picks.append(item)
    return picks


def run_analysis(parsed, target_role, experience_years):
    data = load_data()
    role_data = data["roles"].get(target_role)

    if not role_data:
        return None

    gap = skill_gap(parsed, role_data)

    return {
        "ats": ats_score(parsed, role_data, experience_years),
        "gap": gap,
        "courses": suggest_courses(gap, data),
    }